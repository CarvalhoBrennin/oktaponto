"""Parser do Espelho de Ponto (Excel ou PDF) exportado pelo sistema de ponto.

Diferente do fluxo por OCR, aqui a fonte já é uma tabela estruturada
("Data", "Hora Trab.", "Obs" etc.), então não há incerteza de leitura de
caracteres — o parser só precisa localizar as colunas certas e aplicar as
mesmas regras de inclusão/ignorar usadas para os prints.
"""

from __future__ import annotations

import math
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

from ..models import TimesheetRecord
from .normalizer import expected_weekday, normalize_text
from .timeparse import DATE_RE, parse_hhmm_token

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".pdf"}


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    # bool é subclasse de int em Python. Tratá-lo como número faria True
    # virar 24:00 e False virar 0:00, mascarando uma célula inválida.
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, time):
        return f"{value.hour}:{value.minute:02d}"
    if isinstance(value, timedelta):
        total_minutes = round(value.total_seconds() / 60)
        sign = "-" if total_minutes < 0 else ""
        hours, minutes = divmod(abs(total_minutes), 60)
        return f"{sign}{hours}:{minutes:02d}"
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        try:
            numeric = float(value)
        except (TypeError, ValueError, OverflowError):
            return ""
        if not math.isfinite(numeric):
            return ""
        total_minutes = round(numeric * 24 * 60)
        h, m = divmod(total_minutes, 60)
        return f"{h}:{m:02d}"
    return str(value).strip()


def _classify_header_cell(normalized: str) -> str | None:
    n = normalized.replace(".", "").strip()
    if not n:
        return None
    if n == "DATA":
        return "data"
    if n.startswith("HORA TRAB") or n.startswith("HORAS TRAB") or n == "TRABALHOU":
        return "hora_trab"
    if n in {"OBS", "OBSERVACAO", "OBSERVACOES"} or n.startswith("OBSERVA"):
        return "obs"
    if "ENTRADA" in n:
        if re.search(r"(?:^|\s)2(?:\s|$)", n):
            return "entrada2"
        if re.search(r"(?:^|\s)1(?:\s|$)", n):
            return "entrada1"
        return "entrada1"
    if "SAIDA" in n:
        if re.search(r"(?:^|\s)2(?:\s|$)", n):
            return "saida2"
        if re.search(r"(?:^|\s)1(?:\s|$)", n):
            return "saida1"
        return "saida1"
    return None


def _classify_header_row(cells: list[str]) -> dict[str, int] | None:
    colmap: dict[str, int] = {}
    for idx, raw in enumerate(cells):
        role = _classify_header_cell(normalize_text(raw))
        if role and role not in colmap:
            colmap[role] = idx
    if "data" in colmap and "hora_trab" in colmap:
        return colmap
    return None


def _col(cells: list[str], colmap: dict[str, int], role: str) -> str:
    idx = colmap.get(role)
    if idx is None or idx >= len(cells):
        return ""
    return (cells[idx] or "").strip()


def _row_to_record(cells: list[str], colmap: dict[str, int], source_file: str) -> TimesheetRecord | None:
    date_text = _col(cells, colmap, "data")
    match = DATE_RE.search(date_text)
    if not match:
        return None

    day, month, year = int(match.group(1)), int(match.group(2)), int(match.group(3))
    try:
        record_date = date(year, month, day)
    except ValueError:
        return None

    weekday = expected_weekday(record_date)

    hora_text = _col(cells, colmap, "hora_trab")
    worked_minutes = parse_hhmm_token(hora_text, max_hours=24)

    obs_text = _col(cells, colmap, "obs")
    normalized_obs = normalize_text(obs_text)

    pontos = [
        p for p in (
            _col(cells, colmap, "entrada1"),
            _col(cells, colmap, "saida1"),
            _col(cells, colmap, "entrada2"),
            _col(cells, colmap, "saida2"),
        ) if p
    ]

    notes: list[str] = []
    status = "ok"
    include = True
    tipo = ""

    if worked_minutes is None:
        status = "revisar"
        include = False
        notes.append("Não foi possível localizar o valor de 'Hora Trab.' no espelho de ponto.")
    elif worked_minutes > 24 * 60:
        status = "revisar"
        include = False
        notes.append("Valor de 'Hora Trab.' acima de 24:00; confira manualmente.")
    elif worked_minutes <= 0:
        include = False
        if "FOLGA" in normalized_obs:
            status = "ignorado"
            tipo = "Folga"
            notes.append("Folga com 0:00 ignorada.")
        elif "FERIADO" in normalized_obs:
            status = "ignorado"
            tipo = "Feriado"
            notes.append("Feriado com 0:00 ignorado.")
        elif "NAO REGISTROU" in normalized_obs or pontos:
            status = "revisar"
            notes.append("Dia sem saída registrada no espelho de ponto; confira manualmente.")
        else:
            status = "ignorado"
            notes.append("Registro 0:00 ignorado.")
    else:
        tipo = "Expediente (ajuste aprovado)" if "AJUSTE APROVADO" in normalized_obs else "Expediente"

    cleaned_obs = obs_text.strip(" ()")
    if cleaned_obs and "FOLGA" not in normalized_obs and "FERIADO" not in normalized_obs:
        notes.append(cleaned_obs)

    return TimesheetRecord(
        data=record_date,
        dia_semana=weekday,
        worked_minutes=worked_minutes,
        include=include,
        status=status,
        source_file=source_file,
        tipo=tipo,
        pontos=pontos[:4],
        notes=notes,
        confidence=None,
    )


def parse_espelho_xlsx(path: str | Path) -> list[TimesheetRecord]:
    from openpyxl import load_workbook

    path = Path(path)
    source_file = str(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        all_records: list[TimesheetRecord] = []
        for ws in wb.worksheets:
            colmap: dict[str, int] | None = None
            records: list[TimesheetRecord] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_cell_text(v) for v in row]
                if not any(c for c in cells):
                    continue
                if colmap is None:
                    colmap = _classify_header_row(cells)
                    continue
                candidate_header = _classify_header_row(cells)
                if candidate_header:
                    colmap = candidate_header
                    continue
                record = _row_to_record(cells, colmap, source_file)
                if record:
                    records.append(record)
            if records:
                all_records.extend(records)
        return all_records
    finally:
        wb.close()


def parse_espelho_pdf(path: str | Path) -> list[TimesheetRecord]:
    import pdfplumber

    path = Path(path)
    source_file = str(path)
    colmap: dict[str, int] | None = None
    records: list[TimesheetRecord] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for raw_row in table:
                    if not raw_row:
                        continue
                    cells = [_cell_text(c) for c in raw_row]
                    if not any(c for c in cells):
                        continue
                    if colmap is None:
                        colmap = _classify_header_row(cells)
                        continue
                    candidate_header = _classify_header_row(cells)
                    if candidate_header:
                        colmap = candidate_header
                        continue
                    record = _row_to_record(cells, colmap, source_file)
                    if record:
                        records.append(record)
    return records


def parse_espelho_file(path: str | Path) -> list[TimesheetRecord]:
    """Lê um espelho de ponto em .xlsx/.xlsm ou .pdf e devolve os registros."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return parse_espelho_xlsx(path)
    if suffix == ".pdf":
        return parse_espelho_pdf(path)
    if suffix == ".xls":
        raise ValueError(
            "Formato .xls (Excel antigo) não é suportado. Abra o arquivo no Excel e salve como .xlsx."
        )
    raise ValueError(f"Formato não suportado: {suffix or '(sem extensão)'}. Use .xlsx ou .pdf.")
