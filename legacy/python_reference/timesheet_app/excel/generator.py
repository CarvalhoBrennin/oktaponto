from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, date
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..constants import default_filename_for_month, dia_semana_pt, mes_ano_pt, month_bounds
from ..automation import group_records_by_month, reserve_available_path
from ..file_utils import atomic_write_via_path
from ..models import ProfileConfig, TimesheetRecord


@dataclass(slots=True)
class TimesheetGenerationResult:
    output_path: Path
    included_days: int
    total_minutes: int
    period_start: date
    period_end: date
    mes_ano: str

    @property
    def total_hhmm(self) -> str:
        h, m = divmod(self.total_minutes, 60)
        return f"{h}:{m:02d}"


# Paleta do layout de referência.
C_AZUL_ESCURO = "1F3864"
C_AZUL_MEDIO = "2E75B6"
C_INFO_BG = "EEF4FB"
C_BRANCO = "FFFFFF"
C_SEXTA = "FFF2CC"
C_VERDE = "1E7145"
C_VERMELHO = "C00000"
C_AZUL_TEXTO = "1F3864"
C_BORDA = "C8D8EC"
C_BRANCO_TXT = "FFFFFF"
FONTE = "Segoe UI"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(color: str | None = None, bold: bool = False, size: int = 11, italic: bool = False) -> Font:
    kwargs = dict(name=FONTE, size=size, bold=bold, italic=italic)
    if color:
        kwargs["color"] = color
    return Font(**kwargs)


def _border(color: str = C_BORDA) -> Border:
    side = Side(border_style="thin", color=color)
    return Border(top=side, bottom=side, left=side, right=side)


def _align(horizontal: str = "left", vertical: str = "center") -> Alignment:
    return Alignment(horizontal=horizontal, vertical=vertical)


def _set_safe_value(cell, value) -> None:
    """Atribui texto do usuário sem fórmulas ou caracteres inválidos no XLSX."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub(" ", value)
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def _cell(ws, row: int, col: int, value=None, bg: str | None = None, fg: str | None = None,
          bold: bool = False, size: int = 11, fmt: str | None = None,
          halign: str = "left", italic: bool = False, use_border: bool = True):
    c = ws.cell(row=row, column=col)
    if value is not None:
        _set_safe_value(c, value)
    if bg:
        c.fill = _fill(bg)
    c.font = _font(fg, bold, size, italic)
    c.alignment = _align(halign, "center")
    if fmt:
        c.number_format = fmt
    if use_border:
        c.border = _border()
    return c


def _is_friday_or_weekend(d: date) -> bool:
    return d.weekday() >= 4


def _time_float_from_minutes(minutes: int) -> float:
    return float(minutes * 60) / 86400.0


def _validate_checked_integrity(items: list[TimesheetRecord]) -> None:
    """Falha fechado para qualquer linha marcada que não possa ser exportada."""
    over_24h = [
        record.data for record in items
        if record.include
        and isinstance(record.worked_minutes, int)
        and not isinstance(record.worked_minutes, bool)
        and record.worked_minutes > 24 * 60
    ]
    if over_24h:
        formatted = ", ".join(item.strftime("%d/%m/%Y") for item in sorted(set(over_24h)))
        raise ValueError(f"Há registros acima de 24:00 e eles não podem ser exportados: {formatted}.")

    invalid_included = sorted({record.data for record in items if record.include and not record.eligible_for_excel()})
    if invalid_included:
        formatted = ", ".join(item.strftime("%d/%m/%Y") for item in invalid_included)
        raise ValueError(
            "Há registros marcados para inclusão que ainda são inválidos ou não resolvidos: "
            + formatted + "."
        )


def _validate_records(records: Iterable[TimesheetRecord]) -> list[TimesheetRecord]:
    items = list(records)
    _validate_checked_integrity(items)

    selected = [r for r in items if r.eligible_for_excel()]
    selected.sort(key=lambda item: item.data)
    if not selected:
        raise ValueError("Nenhum registro trabalhado foi selecionado para gerar o timesheet.")
    months = {(r.data.year, r.data.month) for r in selected}
    if len(months) > 1:
        raise ValueError("Os registros selecionados pertencem a mais de um mês. Gere um arquivo por mês.")
    seen_dates: set[date] = set()
    duplicate_dates: set[date] = set()
    for record in selected:
        if record.data in seen_dates:
            duplicate_dates.add(record.data)
        seen_dates.add(record.data)
    if duplicate_dates:
        formatted = ", ".join(item.strftime("%d/%m/%Y") for item in sorted(duplicate_dates))
        raise ValueError(f"Há datas duplicadas entre os registros incluídos: {formatted}.")
    return selected


def generate_timesheet_xlsx(records: Iterable[TimesheetRecord], output_path: str | Path | None = None,
                            profile: ProfileConfig | None = None) -> TimesheetGenerationResult:
    selected = _validate_records(records)
    profile = profile or ProfileConfig()
    first_date = selected[0].data
    period_start, period_end = month_bounds(first_date)
    mes_ano = mes_ano_pt(first_date)
    if output_path is None:
        output_path = Path.cwd() / default_filename_for_month(first_date)
    else:
        output_path = Path(output_path).expanduser()
    if output_path.suffix.casefold() != ".xlsx":
        raise ValueError("O arquivo de saída deve usar a extensão .xlsx.")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Planilha1"
    ws.sheet_view.showGridLines = False

    widths = [13, 18, 17, 25, 22, 22, 52]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # Linha 1: banner.
    ws.row_dimensions[1].height = 36
    ws.merge_cells("A1:G1")
    title_cell = ws.cell(1, 1, f"⏱  TIMESHEET — {mes_ano}")
    title_cell.fill = _fill(C_AZUL_ESCURO)
    title_cell.font = _font(C_BRANCO_TXT, bold=True, size=16)
    title_cell.alignment = _align("center", "center")

    infos = [
        ("COLABORADOR", profile.colaborador),
        ("CARGO", profile.cargo),
        ("PERÍODO", f"{period_start.strftime('%d/%m/%Y')}  →  {period_end.strftime('%d/%m/%Y')}"),
    ]
    for row, (label, value) in enumerate(infos, start=2):
        ws.row_dimensions[row].height = 18
        ws.merge_cells(f"A{row}:B{row}")
        label_cell = ws.cell(row, 1, label)
        label_cell.fill = _fill(C_INFO_BG)
        label_cell.font = _font(C_AZUL_MEDIO, bold=True, size=9)
        label_cell.alignment = _align("left", "center")
        label_cell.border = Border(left=Side(border_style="medium", color=C_AZUL_MEDIO))

        ws.merge_cells(f"C{row}:G{row}")
        value_cell = ws.cell(row, 3)
        _set_safe_value(value_cell, value)
        value_cell.fill = _fill(C_INFO_BG)
        value_cell.font = _font(C_AZUL_TEXTO, size=10)
        value_cell.alignment = _align("left", "center")

    ws.row_dimensions[5].height = 8
    for col in range(1, 8):
        ws.cell(5, col).fill = _fill(C_BRANCO)

    ws.row_dimensions[6].height = 28
    headers = [
        "Data", "Dia da Semana", "Horas Trabalhadas",
        "Colaborador", "Cargo", "Modalidade", "Atividades / Entregas",
    ]
    for col, text in enumerate(headers, 1):
        cell = ws.cell(6, col, text)
        cell.fill = _fill(C_AZUL_MEDIO)
        cell.font = _font(C_BRANCO_TXT, bold=True, size=10)
        cell.alignment = _align("center", "center")
        cell.border = _border(color="1A5C9E")

    first_data_row = 7
    for i, record in enumerate(selected):
        row = first_data_row + i
        ws.row_dimensions[row].height = 18
        if _is_friday_or_weekend(record.data):
            bg = C_SEXTA
        elif i % 2 == 0:
            bg = C_INFO_BG
        else:
            bg = C_BRANCO

        minutes = record.worked_minutes or 0
        if minutes > 480:
            hours_color = C_VERDE
        elif minutes < 480:
            hours_color = C_VERMELHO
        else:
            hours_color = C_AZUL_TEXTO

        time_float = _time_float_from_minutes(minutes)

        _cell(ws, row, 1, datetime(record.data.year, record.data.month, record.data.day),
              bg=bg, fg=C_AZUL_TEXTO, size=10, fmt="DD/MM/YYYY", halign="center")
        _cell(ws, row, 2, dia_semana_pt(record.data), bg=bg, fg=C_AZUL_TEXTO, size=10)
        _cell(ws, row, 3, time_float, bg=bg, fg=hours_color, bold=True, size=10, fmt="[h]:mm", halign="center")
        _cell(ws, row, 4, profile.colaborador, bg=bg, fg=C_AZUL_TEXTO, size=10)
        _cell(ws, row, 5, profile.cargo, bg=bg, fg=C_AZUL_TEXTO, size=10)
        _cell(ws, row, 6, profile.modalidade, bg=bg, fg=C_AZUL_TEXTO, size=10, halign="center")
        activity = (record.atividade or "").strip() or (profile.atividade_padrao or "").strip()
        _cell(ws, row, 7, activity or None, bg=bg)

    last_data_row = first_data_row + len(selected) - 1
    separator_row = last_data_row + 1
    total_row = last_data_row + 2

    ws.row_dimensions[separator_row].height = 6
    for col in range(1, 8):
        ws.cell(separator_row, col).fill = _fill(C_BRANCO)

    total_minutes = sum(record.worked_minutes or 0 for record in selected)
    total_float = _time_float_from_minutes(total_minutes)

    ws.row_dimensions[total_row].height = 28
    ws.merge_cells(f"A{total_row}:B{total_row}")
    total_label_cell = ws.cell(total_row, 1, "TOTAL DE HORAS")
    total_label_cell.fill = _fill(C_AZUL_ESCURO)
    total_label_cell.font = _font(C_BRANCO_TXT, bold=True, size=11)
    total_label_cell.alignment = _align("center", "center")
    total_label_cell.border = _border(color="1A5C9E")

    total_value_cell = ws.cell(total_row, 3, total_float)
    total_value_cell.number_format = "[h]:mm"
    total_value_cell.fill = _fill(C_AZUL_ESCURO)
    total_value_cell.font = _font(C_BRANCO_TXT, bold=True, size=14)
    total_value_cell.alignment = _align("center", "center")
    total_value_cell.border = _border(color="1A5C9E")

    for col in range(4, 8):
        cell = ws.cell(total_row, col)
        cell.fill = _fill(C_AZUL_ESCURO)
        cell.border = _border(color="1A5C9E")

    legend_row = total_row + 2
    ws.row_dimensions[legend_row].height = 14
    ws.merge_cells(f"A{legend_row}:G{legend_row}")
    legend_cell = ws.cell(legend_row, 1, "🟢 > 8h  |  🔴 < 8h  |  🔵 = 8h  |  🟡 Sexta / fim de semana")
    legend_cell.font = _font("888888", size=8, italic=True)
    legend_cell.alignment = _align("left", "center")

    ws.freeze_panes = "A7"
    try:
        atomic_write_via_path(output_path, wb.save)
    finally:
        wb.close()

    return TimesheetGenerationResult(
        output_path=output_path,
        included_days=len(selected),
        total_minutes=total_minutes,
        period_start=period_start,
        period_end=period_end,
        mes_ano=mes_ano,
    )


def generate_timesheet_batch(
    records: Iterable[TimesheetRecord],
    output_dir: str | Path,
    profile: ProfileConfig | None = None,
) -> list[TimesheetGenerationResult]:
    """Gera automaticamente um arquivo por mês presente nos registros incluídos."""
    items = list(records)
    # Valide o conjunto original antes de agrupá-lo. Caso contrário, registros
    # pendentes (que não são elegíveis) desaparecem do agrupamento e um lote
    # poderia ser gerado silenciosamente sem dias que ainda exigiam decisão.
    _validate_checked_integrity(items)
    grouped = group_records_by_month(items)
    if not grouped:
        raise ValueError("Nenhum registro trabalhado foi selecionado para gerar o timesheet.")

    output_dir = Path(output_dir).expanduser()
    output_dir.mkdir(parents=True, exist_ok=True)
    results: list[TimesheetGenerationResult] = []
    created_paths: list[Path] = []
    try:
        for month_records in grouped.values():
            output_path = reserve_available_path(output_dir / default_filename_for_month(month_records[0].data))
            # Registre a reserva antes de iniciar a geração. Assim, mesmo um
            # KeyboardInterrupt imediatamente após o save do XLSX entra no
            # rollback externo. Se a geração falhar, a limpeza abaixo pode
            # tentar remover o mesmo caminho novamente com segurança.
            created_paths.append(output_path)
            try:
                result = generate_timesheet_xlsx(month_records, output_path, profile=profile)
            except BaseException:
                # A reserva existe antes da geração. Se este mês falhar, não
                # deixe um XLSX vazio ocupando o nome para a próxima tentativa.
                try:
                    output_path.unlink(missing_ok=True)
                except OSError:
                    pass
                raise
            results.append(result)
    except BaseException:
        # Uma exportação de vários meses é tratada como uma operação única:
        # se algum mês falhar, removemos os arquivos criados nesta tentativa.
        for created in created_paths:
            try:
                created.unlink(missing_ok=True)
            except OSError:
                pass
        raise
    return results
