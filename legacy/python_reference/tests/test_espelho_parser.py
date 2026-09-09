from __future__ import annotations

import tempfile
import unittest
from datetime import timedelta
from unittest.mock import patch, MagicMock
from pathlib import Path

from openpyxl import Workbook

from timesheet_app.parser.banco_horas_parser import consolidate_records
from timesheet_app.parser.espelho_parser import (
    _cell_text,
    _classify_header_row,
    _row_to_record,
    parse_espelho_file,
    parse_espelho_pdf,
    parse_espelho_xlsx,
)

HEADER = [
    "Data", "1ª Entrada", "1ª Saída", "2ª Entrada", "2ª Saída",
    "Crédito", "Débito", "H. Intervalo", "Hora Trab.", "Saldo", "Obs",
]

ROWS = [
    ["Qua, 01/07/2026", "09:00", "12:53", "14:03", "18:03", "0:00", "0:00", "1:10", "7:53", "87:40", " (ajuste aprovado)"],
    ["Qui, 02/07/2026", "09:42", "16:48", "", "", "0:00", "0:54", "0:00", "7:06", "86:46", ""],
    ["Sáb, 04/07/2026", "", "", "", "", "0:00", "0:00", "0:00", "0:00", "85:52", "Folga"],
    ["Qui, 09/07/2026", "", "", "", "", "0:00", "0:00", "0:00", "0:00", "86:21", "Feriado"],
    ["Seg, 27/07/2026", "12:55", "14:04", "18:06", "", "0:00", "0:00", "4:02", "0:00", "85:56", " (não registrou a saída)"],
]


def _build_sample_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "PontoSimples"
    ws.append(["Fulano de Tal"])
    ws.append(["Cargo de teste"])
    ws.append([])
    ws.append(HEADER)
    for row in ROWS:
        ws.append(row)
    wb.save(path)


class EspelhoCellTextTest(unittest.TestCase):
    def test_excel_timedelta_is_formatted_as_hours_minutes(self) -> None:
        self.assertEqual(_cell_text(timedelta(hours=7, minutes=35)), "7:35")

    def test_negative_timedelta_keeps_sign(self) -> None:
        self.assertEqual(_cell_text(timedelta(minutes=-45)), "-0:45")

    def test_boolean_cell_is_not_misread_as_excel_day_fraction(self) -> None:
        self.assertEqual(_cell_text(True), "True")
        self.assertEqual(_cell_text(False), "False")

    def test_nonfinite_numeric_cells_are_ignored_safely(self) -> None:
        self.assertEqual(_cell_text(float("nan")), "")
        self.assertEqual(_cell_text(float("inf")), "")
        self.assertEqual(_cell_text(float("-inf")), "")


class EspelhoHeaderClassifierTest(unittest.TestCase):
    def test_classify_header_row(self) -> None:
        colmap = _classify_header_row(HEADER)
        self.assertIsNotNone(colmap)
        self.assertEqual(colmap["data"], 0)
        self.assertEqual(colmap["hora_trab"], 8)
        self.assertEqual(colmap["obs"], 10)
        self.assertEqual(colmap["entrada1"], 1)
        self.assertEqual(colmap["saida1"], 2)
        self.assertEqual(colmap["entrada2"], 3)
        self.assertEqual(colmap["saida2"], 4)

    def test_classify_header_with_number_after_label(self) -> None:
        header = ["Data", "Entrada 1", "Saída 1", "Entrada 2", "Saída 2", "Hora Trab."]
        colmap = _classify_header_row(header)
        self.assertEqual(colmap["entrada1"], 1)
        self.assertEqual(colmap["saida1"], 2)
        self.assertEqual(colmap["entrada2"], 3)
        self.assertEqual(colmap["saida2"], 4)

    def test_non_header_row_returns_none(self) -> None:
        self.assertIsNone(_classify_header_row(["Fulano de Tal"]))


class EspelhoRowToRecordTest(unittest.TestCase):
    def setUp(self) -> None:
        self.colmap = _classify_header_row(HEADER)

    def test_normal_day(self) -> None:
        record = _row_to_record(ROWS[1], self.colmap, "espelho.xlsx")
        self.assertEqual(record.data.isoformat(), "2026-07-02")
        self.assertEqual(record.worked_minutes, 7 * 60 + 6)
        self.assertEqual(record.status, "ok")
        self.assertTrue(record.include)
        self.assertTrue(record.eligible_for_excel())
        self.assertEqual(record.pontos, ["09:42", "16:48"])

    def test_ajuste_aprovado_note(self) -> None:
        record = _row_to_record(ROWS[0], self.colmap, "espelho.xlsx")
        self.assertEqual(record.worked_minutes, 7 * 60 + 53)
        self.assertIn("ajuste aprovado", record.notes[0])
        self.assertEqual(record.tipo, "Expediente (ajuste aprovado)")

    def test_folga_is_ignored(self) -> None:
        record = _row_to_record(ROWS[2], self.colmap, "espelho.xlsx")
        self.assertEqual(record.status, "ignorado")
        self.assertFalse(record.include)
        self.assertFalse(record.eligible_for_excel())

    def test_feriado_is_ignored(self) -> None:
        record = _row_to_record(ROWS[3], self.colmap, "espelho.xlsx")
        self.assertEqual(record.status, "ignorado")
        self.assertEqual(record.tipo, "Feriado")

    def test_worked_above_24_hours_is_flagged_for_review(self) -> None:
        row = list(ROWS[1])
        row[8] = "24:30"
        record = _row_to_record(row, self.colmap, "espelho.xlsx")
        self.assertEqual(record.status, "revisar")
        self.assertFalse(record.include)

    def test_missing_checkout_flagged_for_review(self) -> None:
        record = _row_to_record(ROWS[4], self.colmap, "espelho.xlsx")
        self.assertEqual(record.status, "revisar")
        self.assertFalse(record.include)
        self.assertTrue(any("confira manualmente" in n for n in record.notes))


class EspelhoXlsxIntegrationTest(unittest.TestCase):
    def test_parse_sample_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "espelho.xlsx"
            _build_sample_xlsx(path)

            records = parse_espelho_xlsx(path)
            self.assertEqual(len(records), len(ROWS))

            by_date = {r.data.isoformat(): r for r in records}
            self.assertEqual(by_date["2026-07-01"].worked_minutes, 7 * 60 + 53)
            self.assertFalse(by_date["2026-07-04"].include)
            self.assertFalse(by_date["2026-07-09"].include)
            self.assertEqual(by_date["2026-07-27"].status, "revisar")

            consolidated = consolidate_records(records)
            eligible = [r for r in consolidated if r.eligible_for_excel()]
            self.assertEqual(len(eligible), 2)

    def test_parse_collects_records_from_multiple_worksheets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "espelho_multiplas_abas.xlsx"
            wb = Workbook()
            ws1 = wb.active
            ws1.title = "Junho"
            ws1.append(HEADER)
            ws1.append(["Seg, 08/06/2026", "09:00", "12:00", "13:00", "18:00", "", "", "", "8:00", "", ""])
            ws2 = wb.create_sheet("Julho")
            ws2.append(HEADER)
            ws2.append(["Qua, 01/07/2026", "09:00", "12:00", "13:00", "17:30", "", "", "", "7:30", "", ""])
            wb.save(path)

            records = parse_espelho_xlsx(path)
            self.assertEqual([record.data.isoformat() for record in records], ["2026-06-08", "2026-07-01"])


    def test_repeated_header_can_change_column_order_within_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "espelho_duas_tabelas.xlsx"
            wb = Workbook()
            ws = wb.active
            ws.append(["Data", "Hora Trab.", "Obs"])
            ws.append(["Seg, 08/06/2026", "8:00", ""])
            ws.append([])
            ws.append(["Hora Trab.", "Obs", "Data"])
            ws.append(["7:30", "", "Qua, 01/07/2026"])
            wb.save(path)

            records = parse_espelho_xlsx(path)
            self.assertEqual(
                [(record.data.isoformat(), record.worked_minutes) for record in records],
                [("2026-06-08", 480), ("2026-07-01", 450)],
            )

    def test_parse_espelho_file_dispatches_by_extension(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "espelho.xlsx"
            _build_sample_xlsx(path)
            records = parse_espelho_file(path)
            self.assertEqual(len(records), len(ROWS))


    def test_pdf_collects_all_tables_from_a_page(self) -> None:
        page = MagicMock()
        page.extract_tables.return_value = [
            [["Data", "Hora Trab."], None, [], ["Seg, 08/06/2026", "8:00"]],
            [["Data", "Hora Trab."], ["Qua, 01/07/2026", "7:30"]],
        ]
        pdf = MagicMock()
        pdf.pages = [page]
        context = MagicMock()
        context.__enter__.return_value = pdf
        context.__exit__.return_value = False
        with patch("pdfplumber.open", return_value=context):
            records = parse_espelho_pdf("espelho.pdf")
        self.assertEqual(
            [(record.data.isoformat(), record.worked_minutes) for record in records],
            [("2026-06-08", 480), ("2026-07-01", 450)],
        )

    def test_unsupported_extension_raises(self) -> None:
        with self.assertRaises(ValueError):
            parse_espelho_file("espelho.docx")

    def test_old_xls_extension_raises_clear_message(self) -> None:
        with self.assertRaises(ValueError):
            parse_espelho_file("espelho.xls")


if __name__ == "__main__":
    unittest.main()
