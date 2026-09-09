from __future__ import annotations

import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from timesheet_app.excel.generator import generate_timesheet_batch, generate_timesheet_xlsx
from timesheet_app.models import ProfileConfig, TimesheetRecord


class ExcelGeneratorTest(unittest.TestCase):
    def test_generate_xlsx_with_numeric_total(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 8 * 60 + 30, True),
            TimesheetRecord(date(2026, 6, 9), "Terça-feira", 7 * 60 + 9, True),
            TimesheetRecord(date(2026, 6, 7), "Domingo", 0, False),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet_junho_2026.xlsx"
            result = generate_timesheet_xlsx(records, output, profile=ProfileConfig())
            self.assertTrue(output.exists())
            self.assertEqual(result.total_minutes, 15 * 60 + 39)

            wb = load_workbook(output, data_only=False)
            ws = wb["Planilha1"]
            self.assertFalse(ws.sheet_view.showGridLines)
            self.assertEqual(ws.freeze_panes, "A7")
            # Dados começam na linha 7; Domingo 0:00 não entra.
            self.assertEqual(ws["A7"].value.strftime("%d/%m/%Y"), "08/06/2026")
            self.assertEqual(ws["C7"].number_format, "[h]:mm")
            total_cell = ws["C10"]
            self.assertEqual(total_cell.number_format, "[h]:mm")
            self.assertFalse(isinstance(total_cell.value, str) and total_cell.value.startswith("="))
            self.assertNotEqual(total_cell.data_type, "f")
            # openpyxl pode reler células [h]:mm como timedelta, mas o arquivo não deve conter fórmula.
            if hasattr(total_cell.value, "total_seconds"):
                self.assertEqual(int(total_cell.value.total_seconds() // 60), 15 * 60 + 39)
            else:
                self.assertAlmostEqual(float(total_cell.value) * 24 * 60, 15 * 60 + 39, places=3)



    def test_daily_24_hours_is_displayed_as_24_not_midnight(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 24 * 60, True, "ok")]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            generate_timesheet_xlsx(records, output, profile=ProfileConfig())
            wb = load_workbook(output, data_only=True)
            cell = wb["Planilha1"]["C7"]
            self.assertEqual(cell.number_format, "[h]:mm")
            self.assertTrue(hasattr(cell.value, "total_seconds"))
            self.assertEqual(int(cell.value.total_seconds() // 60), 24 * 60)
            wb.close()

    def test_activity_uses_record_override_then_profile_default(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True, atividade="Deploy"),
            TimesheetRecord(date(2026, 6, 9), "Terça-feira", 480, True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            generate_timesheet_xlsx(
                records, output,
                profile=ProfileConfig(atividade_padrao="Desenvolvimento"),
            )
            wb = load_workbook(output, data_only=True)
            ws = wb["Planilha1"]
            self.assertEqual(ws["G7"].value, "Deploy")
            self.assertEqual(ws["G8"].value, "Desenvolvimento")

    def test_batch_generates_one_file_per_month_and_does_not_overwrite(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True),
            TimesheetRecord(date(2026, 7, 1), "Quarta-feira", 450, True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            first = generate_timesheet_batch(records, tmp, profile=ProfileConfig())
            second = generate_timesheet_batch(records, tmp, profile=ProfileConfig())
            self.assertEqual(len(first), 2)
            self.assertEqual(len(second), 2)
            self.assertTrue(all(result.output_path.exists() for result in first + second))
            self.assertNotEqual(first[0].output_path, second[0].output_path)

    def test_edited_record_respects_include_false(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 9), "Terça-feira", 7 * 60, True, "ok"),
            TimesheetRecord(date(2026, 6, 10), "Quarta-feira", 8 * 60, False, "editado"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet_junho_2026.xlsx"
            result = generate_timesheet_xlsx(records, output, profile=ProfileConfig())
            self.assertEqual(result.included_days, 1)
            self.assertEqual(result.total_minutes, 7 * 60)

    def test_user_text_starting_with_equals_is_not_written_as_formula(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True, atividade="=HYPERLINK(\"x\")"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            generate_timesheet_xlsx(
                records,
                output,
                profile=ProfileConfig(colaborador="=1+1"),
            )
            wb = load_workbook(output, data_only=False)
            ws = wb["Planilha1"]
            self.assertEqual(ws["C2"].data_type, "s")
            self.assertEqual(ws["D7"].data_type, "s")
            self.assertEqual(ws["G7"].data_type, "s")
            self.assertEqual(ws["C2"].value, "=1+1")
            self.assertEqual(ws["D7"].value, "=1+1")

    def test_record_above_24_hours_is_rejected_at_export_boundary(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 24 * 60 + 1, True)]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "acima de 24:00"):
                generate_timesheet_xlsx(records, Path(tmp) / "timesheet.xlsx", profile=ProfileConfig())

    def test_duplicate_included_dates_are_rejected(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True),
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 450, True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "datas duplicadas"):
                generate_timesheet_xlsx(records, Path(tmp) / "timesheet.xlsx", profile=ProfileConfig())

    def test_export_derives_weekday_from_date(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Domingo", 480, True)]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            generate_timesheet_xlsx(records, output, profile=ProfileConfig())
            wb = load_workbook(output, data_only=True)
            self.assertEqual(wb["Planilha1"]["B7"].value, "Segunda-feira")
            wb.close()

    def test_excel_illegal_control_characters_are_sanitized(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True, atividade="Deploy\x00ok")]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            generate_timesheet_xlsx(records, output, profile=ProfileConfig())
            wb = load_workbook(output, data_only=True)
            self.assertEqual(wb["Planilha1"]["G7"].value, "Deploy ok")
            wb.close()

    def test_atomic_save_preserves_existing_file_on_failure(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True)]
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "timesheet.xlsx"
            output.write_bytes(b"arquivo-anterior")

            def broken_save(_self, temp_path) -> None:
                Path(temp_path).write_bytes(b"parcial")
                raise OSError("falha simulada")

            with patch("openpyxl.workbook.workbook.Workbook.save", new=broken_save):
                with self.assertRaises(OSError):
                    generate_timesheet_xlsx(records, output, profile=ProfileConfig())

            self.assertEqual(output.read_bytes(), b"arquivo-anterior")

    def test_batch_rolls_back_on_keyboard_interrupt(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True),
            TimesheetRecord(date(2026, 7, 1), "Quarta-feira", 450, True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            original = generate_timesheet_xlsx
            calls = 0

            def interrupt_second(month_records, output_path, profile=None):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise KeyboardInterrupt
                return original(month_records, output_path, profile=profile)

            with patch("timesheet_app.excel.generator.generate_timesheet_xlsx", side_effect=interrupt_second):
                with self.assertRaises(KeyboardInterrupt):
                    generate_timesheet_batch(records, tmp, profile=ProfileConfig())

            self.assertEqual(list(Path(tmp).glob("*.xlsx")), [])

    def test_batch_tracks_reserved_path_before_generation_returns(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True)]
        with tempfile.TemporaryDirectory() as tmp:
            created = []

            def interrupt_after_writing(month_records, output_path, profile=None):
                Path(output_path).write_bytes(b"xlsx-parcial")
                raise KeyboardInterrupt

            with patch("timesheet_app.excel.generator.generate_timesheet_xlsx", side_effect=interrupt_after_writing):
                with self.assertRaises(KeyboardInterrupt):
                    generate_timesheet_batch(records, tmp, profile=ProfileConfig())

            self.assertEqual(created, [])
            self.assertEqual(list(Path(tmp).glob("*.xlsx")), [])

    def test_batch_rolls_back_files_created_before_later_failure(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True),
            TimesheetRecord(date(2026, 7, 1), "Quarta-feira", 450, True),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            original = generate_timesheet_xlsx
            calls = 0

            def fail_second(month_records, output_path, profile=None):
                nonlocal calls
                calls += 1
                if calls == 2:
                    raise OSError("falha no segundo mês")
                return original(month_records, output_path, profile=profile)

            with patch("timesheet_app.excel.generator.generate_timesheet_xlsx", side_effect=fail_second):
                with self.assertRaises(OSError):
                    generate_timesheet_batch(records, tmp, profile=ProfileConfig())

            self.assertEqual(list(Path(tmp).glob("*.xlsx")), [])

    def test_wrong_output_extension_is_rejected(self) -> None:
        records = [TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True)]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "extensão .xlsx"):
                generate_timesheet_xlsx(records, Path(tmp) / "timesheet.csv")

    def test_batch_does_not_hide_included_value_above_24_hours(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True, "ok"),
            TimesheetRecord(date(2026, 6, 9), "Terça-feira", 24 * 60 + 1, True, "ok"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "acima de 24:00"):
                generate_timesheet_batch(records, tmp)
            self.assertEqual(list(Path(tmp).glob("*.xlsx")), [])

    def test_incoherent_checked_record_is_rejected_instead_of_silently_filtered(self) -> None:
        records = [
            TimesheetRecord(date(2026, 6, 8), "Segunda-feira", 480, True, "ok"),
            TimesheetRecord(date(2026, 6, 9), "Terça-feira", 450, True, "revisar"),
        ]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError, "marcados para inclusão"):
                generate_timesheet_xlsx(records, Path(tmp) / "timesheet.xlsx")
            with self.assertRaisesRegex(ValueError, "marcados para inclusão"):
                generate_timesheet_batch(records, tmp)


if __name__ == "__main__":
    unittest.main()
